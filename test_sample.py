# Optional self-test. Run from this folder after installing requirements:
#     python test_sample.py
# It checks the supplied 23 SEP sample PDF against the 23 SEP required template.
import sys
from pathlib import Path

from app import parse_pdf, group_rotations, make_excel
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
PDF = ROOT / 'sample_input.pdf'
TEMPLATE = ROOT / 'templates' / 'MDV_Schedule_Template.xlsx'

if not PDF.exists():
    print('Copy a representative PDF to sample_input.pdf to run this self-test.')
    sys.exit(1)

date_string, rows = parse_pdf(PDF)
blocks = group_rotations(rows)
out = make_excel(blocks, date_string)

expected = load_workbook(TEMPLATE, data_only=True).active
actual = load_workbook(out, data_only=True).active
mismatches = []
for r in range(6, 34):
    e = [expected.cell(r, c).value for c in range(3, 11)]
    a = [actual.cell(r, c).value for c in range(3, 11)]
    if e != a:
        mismatches.append((r, e, a))

print(f'Parsed rows: {len(rows)}')
print(f'Schedule blocks: {len(blocks)}')
print(f'Output: {out}')
print('PASS' if not mismatches else f'FAIL: {len(mismatches)} mismatches')
for m in mismatches:
    print(m)
